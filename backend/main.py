import os
from io import BytesIO
from typing import Literal

import requests
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

load_dotenv()
app = FastAPI(title="LegalEase API", version="2.0.0")
app.state.ai_provider = os.getenv("AI_PROVIDER", "groq").strip().lower()
allowed_origins = os.getenv("FRONTEND_ORIGIN", "http://localhost:5173,http://127.0.0.1:5173").split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in allowed_origins if origin.strip()],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


class DraftRequest(BaseModel):
    document_type: str = Field(min_length=2, max_length=100)
    party_one: str = Field(min_length=1, max_length=200)
    party_two: str = Field(min_length=1, max_length=200)
    effective_date: str
    key_terms: str = Field(min_length=5, max_length=10000)


class TextRequest(BaseModel):
    text: str = Field(min_length=1, max_length=30000)


class AIProviderRequest(BaseModel):
    provider: Literal["groq", "gemini"]


def ai_provider_settings() -> dict:
    providers = {
        "groq": {
            "model": os.getenv("GROQ_MODEL", "openai/gpt-oss-120b"),
            "configured": bool(os.getenv("GROQ_API_KEY")),
        },
        "gemini": {
            "model": os.getenv("GEMINI_MODEL", "gemini-3.8-flash"),
            "configured": bool(os.getenv("GEMINI_API_KEY")),
        },
    }
    active_provider = app.state.ai_provider
    return {
        "provider": active_provider,
        "model": providers.get(active_provider, {}).get("model", ""),
        "providers": providers,
    }


def ask_ai(prompt: str) -> str:
    provider = app.state.ai_provider
    try:
        if provider == "groq":
            key = os.getenv("GROQ_API_KEY")
            if not key:
                raise HTTPException(503, "Add GROQ_API_KEY to the backend .env file to enable AI drafting.")
            result = requests.post(
                "https://api.groq.com/openai/v1/chat/completions",
                headers={"Authorization": f"Bearer {key}"},
                json={"model": os.getenv("GROQ_MODEL", "openai/gpt-oss-120b"), "messages": [{"role": "user", "content": prompt}], "temperature": 0.35},
                timeout=90,
            )
            result.raise_for_status()
            return result.json()["choices"][0]["message"]["content"].strip()
        if provider == "gemini":
            key = os.getenv("GEMINI_API_KEY")
            if not key:
                raise HTTPException(503, "Add GEMINI_API_KEY to the backend .env file to enable AI drafting.")
            model = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")
            result = requests.post(
                f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent",
                params={"key": key}, json={"contents": [{"parts": [{"text": prompt}]}]}, timeout=90,
            )
            result.raise_for_status()
            return result.json()["candidates"][0]["content"]["parts"][0]["text"].strip()
        raise HTTPException(503, "Unsupported AI_PROVIDER. Choose 'groq' or 'gemini'.")
    except HTTPException:
        raise
    except requests.HTTPError as exc:
        response = exc.response
        status = response.status_code if response is not None else "unknown"
        message = "The provider rejected the request."
        if response is not None:
            try:
                message = response.json().get("error", {}).get("message", message)
            except (ValueError, AttributeError):
                pass
        raise HTTPException(502, f"AI provider returned HTTP {status}: {str(message)[:600]}") from exc
    except requests.RequestException as exc:
        raise HTTPException(502, f"The configured AI provider could not complete the request: {exc}") from exc
    except (KeyError, IndexError, ValueError) as exc:
        raise HTTPException(502, "The AI provider returned an unexpected response.") from exc


@app.get("/")
def health():
    return {"message": "LegalEase API is running", "version": "2.0.0"}


@app.get("/api/settings/ai-provider")
def get_ai_provider():
    return ai_provider_settings()


@app.put("/api/settings/ai-provider")
def set_ai_provider(request: AIProviderRequest):
    key_name = "GROQ_API_KEY" if request.provider == "groq" else "GEMINI_API_KEY"
    if not os.getenv(key_name):
        raise HTTPException(503, f"Configure {key_name} in the backend .env before switching providers.")
    app.state.ai_provider = request.provider
    return ai_provider_settings()


@app.post("/api/documents/generate")
def generate(request: DraftRequest):
    prompt = f"""Draft a clear, structured {request.document_type} based only on the supplied details. Do not invent facts or claim jurisdiction-specific validity. Use a professional title, numbered sections, plain language, and signature blocks. Mark material missing information as [TO BE COMPLETED]. Include a brief informational-draft notice at the end.\n\nParty one: {request.party_one}\nParty two: {request.party_two}\nEffective date: {request.effective_date}\nAgreed terms: {request.key_terms}"""
    return {"document_type": request.document_type, "document": ask_ai(prompt)}


@app.post("/api/ai/summary")
def summary(request: TextRequest):
    return {"result": ask_ai("Summarize this document in 3 concise bullets. Identify key obligations and dates; do not provide legal advice.\n\n" + request.text)}


@app.post("/api/ai/explain")
def explain(request: TextRequest):
    return {"result": ask_ai("Explain this legal clause in plain language, state what each party appears to do, and note any ambiguity. Do not give legal advice.\n\n" + request.text)}


@app.post("/api/exports/{format}")
def export_document(format: Literal["txt", "docx", "pdf"], request: TextRequest):
    filename = "legalease-document"
    if format == "txt":
        data, media = request.text.encode("utf-8"), "text/plain; charset=utf-8"
    elif format == "docx":
        from docx import Document
        doc = Document()
        for index, line in enumerate(request.text.splitlines()):
            if index == 0:
                doc.add_heading(line or "Legal Document", 0)
            elif line.strip():
                doc.add_paragraph(line)
        buffer = BytesIO(); doc.save(buffer); data = buffer.getvalue()
        media = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    else:
        from reportlab.lib.pagesizes import LETTER
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from xml.sax.saxutils import escape
        buffer = BytesIO(); pdf = SimpleDocTemplate(buffer, pagesize=LETTER, rightMargin=64, leftMargin=64, topMargin=64, bottomMargin=64)
        styles = getSampleStyleSheet(); story = []
        for index, line in enumerate(request.text.splitlines()):
            if line.strip(): story.append(Paragraph(escape(line), styles["Title"] if index == 0 else styles["BodyText"])); story.append(Spacer(1, 8))
        pdf.build(story); data = buffer.getvalue(); media = "application/pdf"
    return StreamingResponse(BytesIO(data), media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}.{format}"'})
